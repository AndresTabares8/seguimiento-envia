"""Una ejecución por lote. GitHub Actions administra el horario."""
import html
import json
import os
from pathlib import Path
import re
import time
import unicodedata
from datetime import datetime, timezone
from urllib.parse import urlparse

from adaptador_envia import consultar, ConsultaError


def normalizar(texto):
    texto = unicodedata.normalize('NFKD', texto)
    return ' '.join(''.join(c for c in texto if not unicodedata.combining(c)).upper().split())


# Ampliar este mapa con respuestas actuales comprobadas. Nunca buscar subcadenas.
ESTADOS = {
    'ENTREGADO': 'Entregado',
    'EN REPARTO': 'En reparto',
    'EN DISTRIBUCION': 'En reparto',
    'NOVEDAD': 'Novedad',
    'CON NOVEDAD': 'Novedad',
    'NO ENTREGADO': 'Novedad',
    'EN TRANSITO': 'En tránsito',
    'EN CIUDAD DE ENTREGA': 'En ciudad de entrega',
    'GUIA GENERADA': 'Guía generada',
    'ENTREGAMOS TU ENVIO': 'Entregado',
    '¡PREPARATE PARA RECIBIRLO!': 'En reparto',
    '¡UPS! ALGO PASA CON TU ENVIO': 'Novedad',
    '¡FALTA POCO!': 'En ciudad de entrega',
    'TU ENVIO ESTA EN MOVIMIENTO.': 'En tránsito',
    'TU GUIA FUE GENERADA': 'Guía generada',
}


def clasificar(respuesta, guia):
    if not isinstance(respuesta, dict) or respuesta.get('guia') != guia:
        return 'Sin verificar'
    estado = respuesta.get('estado_actual')
    evidencia = respuesta.get('evidencia_actual')
    fuente = respuesta.get('fuente', '')
    if not all(isinstance(x, str) and x.strip() for x in (estado, evidencia, fuente)):
        return 'Sin verificar'
    url = urlparse(fuente)
    host = url.hostname or ''
    if url.scheme != 'https' or not (host == 'envia.co' or host.endswith('.envia.co')):
        return 'Sin verificar'
    # La evidencia debe ser el campo de estado vigente, no un menú ni el historial.
    if normalizar(estado) != normalizar(evidencia):
        return 'Sin verificar'
    return ESTADOS.get(normalizar(estado), 'Sin verificar')


def validar_guias(datos):
    if not isinstance(datos, list) or len(datos) != 20:
        raise ValueError('Se requieren exactamente 20 guías')
    if any(not isinstance(g, str) or not re.fullmatch(r'[0-9]{5,30}', g) for g in datos):
        raise ValueError('Las guías deben ser texto de 5 a 30 dígitos')
    if len(set(datos)) != 20:
        raise ValueError('Hay guías repetidas')
    return datos


def leer_excel(ruta='guias.xlsx'):
    from openpyxl import load_workbook
    libro = load_workbook(ruta, read_only=True, data_only=False)
    try:
        filas = list(libro.active.iter_rows(values_only=True))
        if not filas:
            raise ValueError('El Excel está vacío')
        cabeceras = [str(v or '').strip() for v in filas[0]]
        nombre = next((n for n in ('Número de guía', 'guia') if n in cabeceras), None)
        if not nombre or cabeceras.count(nombre) != 1:
            raise ValueError('Se requiere una columna Número de guía o guia')
        indice = cabeceras.index(nombre)
        datos = [f[indice].strip() if isinstance(f[indice], str) else f[indice]
                 for f in filas[1:] if f[indice] is not None]
        # Validar antes de quitar la repetición. No convertir números ni fórmulas.
        if any(not isinstance(g, str) or not re.fullmatch(r'[0-9]{5,30}', g) for g in datos):
            raise ValueError('Las guías del Excel deben ser texto con dígitos')
        return validar_guias(list(dict.fromkeys(datos)))
    finally:
        libro.close()


def leer_guias():
    secreto = os.environ.get('GUIAS_JSON')
    if secreto:
        return validar_guias(json.loads(secreto))
    return leer_excel()


def escribir_pagina(filas, fecha):
    # Lista permitida: nunca publicar números, enlaces, respuestas ni texto del proveedor.
    publicas = [{k: fila.get(k, '') for k in ('referencia', 'estado', 'evento_fecha', 'consulta_utc')} for fila in filas]
    carpeta = Path('site')
    carpeta.mkdir(exist_ok=True)
    (carpeta / 'estado.json').write_text(json.dumps({'generado_utc': fecha, 'guias': publicas}, ensure_ascii=False), encoding='utf-8')
    cuerpo = ''.join('<tr><td>' + html.escape(f['referencia']) + '</td><td>' + html.escape(f['estado']) + '</td><td>' + html.escape(f['evento_fecha'] or 'Sin confirmar') + '</td></tr>' for f in publicas)
    pagina = '''<!doctype html><html lang="es"><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<meta name="robots" content="noindex,nofollow"><title>Seguimiento Envía</title>
<style>body{font:16px system-ui;margin:24px auto;padding:0 14px;max-width:850px}td,th{padding:12px 6px;text-align:left;border-bottom:1px solid #ddd}table{width:100%;font-size:14px}#aviso{color:#9b3000}button{padding:12px}</style>
<h1>Seguimiento Envía</h1><p>Reporte de consultas al rastreador oficial. Revisa la fecha antes de usar los estados.</p>
<p id="fecha"></p><p id="aviso"></p><button onclick="location.reload()">Actualizar vista</button>
<table><thead><tr><th>Referencia</th><th>Estado</th><th>Fecha del evento</th></tr></thead><tbody>''' + cuerpo + '''</tbody></table>
<p>Sin verificar indica falta de evidencia o un fallo de consulta. No confirma entrega.</p>
<script>const fecha=new Date(''' + json.dumps(fecha) + ''');function actualizar(){document.getElementById('fecha').textContent='Lote generado: '+fecha.toLocaleString('es-CO',{timeZone:'America/Bogota'})+' (Colombia)';document.getElementById('aviso').textContent=Date.now()-fecha.getTime()>75*60000?'Datos antiguos. Revisa Actions antes de usar estos estados.':''}actualizar();setInterval(actualizar,60000);setTimeout(()=>location.reload(),300000)</script></html>'''
    (carpeta / 'index.html').write_text(pagina, encoding='utf-8')


def main():
    guias = leer_guias()
    fecha = datetime.now(timezone.utc).isoformat()
    filas = []
    detener_lote = False
    for i, guia in enumerate(guias, 1):
        evento_fecha = ''
        try:
            if detener_lote:
                raise ConsultaError('Consultas suspendidas en este lote')
            respuesta = consultar(guia)
            estado = clasificar(respuesta, guia)
            if estado != 'Sin verificar':
                evento_fecha = respuesta.get('evento_fecha', '')
        except ConsultaError as exc:
            estado = 'Sin verificar'
            detener_lote = detener_lote or exc.detener_lote
        except Exception:
            estado = 'Sin verificar'
        filas.append({'referencia': f'G{i:02}', 'estado': estado, 'evento_fecha': evento_fecha,
                      'consulta_utc': datetime.now(timezone.utc).isoformat()})
        if i != len(guias):
            time.sleep(1)
    escribir_pagina(filas, fecha)
    fallos = sum(f['estado'] == 'Sin verificar' for f in filas)
    Path('salud.json').write_text(json.dumps({'sin_verificar': fallos}), encoding='utf-8')
    resumen = os.environ.get('GITHUB_STEP_SUMMARY')
    if resumen:
        with open(resumen, 'a', encoding='utf-8') as archivo:
            archivo.write(f'Lote de {len(guias)} referencias. Sin verificar: {fallos}. Consulta el artefacto o Pages.\n')
    print(f'Lote procesado: {len(guias)} referencias. Sin verificar: {fallos}.')


if __name__ == '__main__':
    main()
